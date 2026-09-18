"""
Part-I: convert every PDF / DOCX / XLSX in data/raw/<Department>/ to Markdown with Docling.

For each source file this writes, under data/converted/<Department>/:
  <name>.md          Markdown export (used for indexing and fidelity checks)
  <name>.json        Lossless DoclingDocument export (layout, tables, page provenance)
and one catalogue for the whole corpus: data/converted/catalog.json

Unchanged files (same SHA-256) are skipped unless --force is given.

Usage:
    python -m src.parser                # convert all departments
    python -m src.parser --dept HR      # one department
    python -m src.parser --force --ocr  # reconvert everything with OCR on
"""
import argparse
import csv
import hashlib
import json
import logging
import time
from datetime import datetime, timezone
from pathlib import Path

from docling.datamodel.base_models import InputFormat
from docling.datamodel.pipeline_options import PdfPipelineOptions, TableFormerMode
from docling.document_converter import DocumentConverter, PdfFormatOption

from src.config import settings

log = logging.getLogger("parser")
SUPPORTED = {".pdf", ".docx", ".xlsx"}
CATALOG = settings.converted_dir / "catalog.json"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1 << 16), b""):
            h.update(block)
    return h.hexdigest()


def load_registry() -> dict:
    """Document registry maintained by each department (title, version, status, access level)."""
    if not settings.registry_csv.exists():
        return {}
    with settings.registry_csv.open(encoding="utf-8") as fh:
        return {row["file"]: row for row in csv.DictReader(fh)}


def build_converter(ocr: bool) -> DocumentConverter:
    pdf_options = PdfPipelineOptions()
    if settings.docling_artifacts_path:           # pre-downloaded models (offline / Docker image)
        pdf_options.artifacts_path = settings.docling_artifacts_path
    pdf_options.do_ocr = ocr                      # digital PDFs do not need OCR; enable for scans
    pdf_options.do_table_structure = True         # TableFormer table recognition
    pdf_options.table_structure_options.mode = TableFormerMode.ACCURATE
    pdf_options.table_structure_options.do_cell_matching = True
    return DocumentConverter(
        allowed_formats=[InputFormat.PDF, InputFormat.DOCX, InputFormat.XLSX],
        format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=pdf_options)},
    )


def to_markdown(doc, src: Path) -> str:
    """Docling Markdown export. For workbooks, Docling maps each sheet to a page; we keep the
    sheet name as a heading so that context is not lost."""
    if src.suffix.lower() != ".xlsx":
        return doc.export_to_markdown()
    from openpyxl import load_workbook
    sheet_names = load_workbook(src, read_only=True).sheetnames
    parts = []
    for page_no in sorted(doc.pages):
        name = sheet_names[page_no - 1] if page_no - 1 < len(sheet_names) else f"Sheet {page_no}"
        parts.append(f"## Sheet: {name}\n\n{doc.export_to_markdown(page_no=page_no)}")
    return "\n\n".join(parts)


def metadata_for(src: Path, registry: dict) -> dict:
    rel = src.relative_to(settings.raw_dir).as_posix()
    department = src.parent.name
    entry = registry.get(rel, {})
    return {
        "doc_id": src.stem,
        "source_file": src.name,
        "source_path": rel,
        "department": entry.get("department", department),
        "doc_type": entry.get("doc_type", "Unknown"),
        "title": entry.get("title", src.stem.replace("_", " ")),
        "version": entry.get("version", ""),
        "effective_date": entry.get("effective_date", ""),
        "status": entry.get("status", "current"),
        "access_level": entry.get("access_level", "internal"),
        "owner": entry.get("owner", ""),
        "file_type": src.suffix.lstrip(".").upper(),
    }


def convert_all(departments=None, force=False, ocr=False) -> list:
    registry = load_registry()
    catalog = {}
    if CATALOG.exists() and not force:
        catalog = {c["source_path"]: c for c in json.loads(CATALOG.read_text())}

    files = sorted(p for p in settings.raw_dir.rglob("*") if p.suffix.lower() in SUPPORTED)
    if departments:
        files = [f for f in files if f.parent.name in departments]
    if not files:
        log.warning("No documents found under %s", settings.raw_dir)
        return []

    converter = build_converter(ocr)
    for src in files:
        meta = metadata_for(src, registry)
        digest = sha256(src)
        out_dir = settings.converted_dir / meta["department"]
        md_path, json_path = out_dir / f"{src.stem}.md", out_dir / f"{src.stem}.json"
        previous = catalog.get(meta["source_path"])
        if previous and previous["sha256"] == digest and md_path.exists():
            log.info("skip (unchanged)  %s", meta["source_path"])
            continue

        out_dir.mkdir(parents=True, exist_ok=True)
        start = time.perf_counter()
        try:
            result = converter.convert(src)
            doc = result.document
            markdown = to_markdown(doc, src)
            md_path.write_text(markdown, encoding="utf-8")
            doc.save_as_json(json_path)
            status, error = str(result.status.value), ""
            stats = {"pages": len(doc.pages), "tables": len(doc.tables), "pictures": len(doc.pictures),
                     "markdown_chars": len(markdown)}
        except Exception as exc:  # keep going; failures are recorded in the catalogue
            status, error, stats = "failure", str(exc), {}
            log.exception("failed  %s", meta["source_path"])

        elapsed = round(time.perf_counter() - start, 2)
        catalog[meta["source_path"]] = {
            **meta, **stats,
            "sha256": digest,
            "markdown_path": md_path.relative_to(settings.converted_dir.parent.parent).as_posix(),
            "conversion_status": status,
            "conversion_seconds": elapsed,
            "error": error,
            "converted_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        }
        log.info("%-8s %6.2fs  %s", status, elapsed, meta["source_path"])

    settings.converted_dir.mkdir(parents=True, exist_ok=True)
    entries = sorted(catalog.values(), key=lambda c: c["source_path"])
    CATALOG.write_text(json.dumps(entries, indent=2), encoding="utf-8")
    ok = sum(1 for c in entries if c["conversion_status"] == "success")
    log.info("catalogue: %d documents, %d converted successfully -> %s", len(entries), ok, CATALOG)
    return entries


def main() -> None:
    ap = argparse.ArgumentParser(description="Convert departmental documents to Markdown with Docling")
    ap.add_argument("--dept", nargs="*", help="Only these department folders")
    ap.add_argument("--force", action="store_true", help="Reconvert even if unchanged")
    ap.add_argument("--ocr", action="store_true", default=settings.docling_ocr, help="Enable OCR for scanned PDFs")
    args = ap.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    convert_all(args.dept, args.force, args.ocr)


if __name__ == "__main__":
    main()
