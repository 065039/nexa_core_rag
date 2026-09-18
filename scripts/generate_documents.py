"""
Generate the synthetic NexaCore document corpus (PDF, DOCX, XLSX) into
data/raw/<Department>/ and write the document registry CSV.

Usage:
    python scripts/generate_documents.py
"""
import csv
import sys
from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (ListFlowable, ListItem, Paragraph, SimpleDocTemplate,
                                Spacer, Table, TableStyle)

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from corpus import DOCUMENTS, ORG  # noqa: E402

RAW = ROOT / "data" / "raw"
REGISTRY = ROOT / "data" / "document_registry.csv"
REGISTRY_FIELDS = ["file", "department", "doc_type", "title", "version",
                   "effective_date", "status", "access_level", "owner"]


def build_pdf(doc: dict, path: Path) -> None:
    styles = getSampleStyleSheet()
    body = styles["BodyText"]
    story = []

    def header_footer(canvas, pdf):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(2 * cm, A4[1] - 1.2 * cm, f"{ORG} | {doc['department'].replace('_', ' ')} | Internal")
        canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Page {pdf.page}")
        canvas.restoreState()

    for kind, content in doc["blocks"]:
        if kind == "h1":
            story += [Paragraph(content, styles["Title"]), Spacer(1, 8)]
        elif kind == "h2":
            story += [Spacer(1, 6), Paragraph(content, styles["Heading2"])]
        elif kind == "p":
            story += [Paragraph(content, body), Spacer(1, 4)]
        elif kind == "bullets":
            story.append(ListFlowable([ListItem(Paragraph(b, body)) for b in content],
                                      bulletType="bullet", leftIndent=14))
        elif kind == "table":
            cells = [[Paragraph(str(c), body) for c in row] for row in content]
            table = Table(cells, repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F1")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story += [table, Spacer(1, 6)]
    SimpleDocTemplate(str(path), pagesize=A4, title=doc["title"], author=ORG,
                      topMargin=2 * cm, bottomMargin=2 * cm).build(
        story, onFirstPage=header_footer, onLaterPages=header_footer)


def build_docx(doc: dict, path: Path) -> None:
    d = DocxDocument()
    d.core_properties.title = doc["title"]
    d.core_properties.author = ORG
    for kind, content in doc["blocks"]:
        if kind == "h1":
            d.add_heading(content, level=1)
        elif kind == "h2":
            d.add_heading(content, level=2)
        elif kind == "p":
            d.add_paragraph(content)
        elif kind == "bullets":
            for b in content:
                d.add_paragraph(b, style="List Bullet")
        elif kind == "table":
            t = d.add_table(rows=len(content), cols=len(content[0]))
            t.style = "Table Grid"
            for r, row in enumerate(content):
                for c, val in enumerate(row):
                    t.cell(r, c).text = str(val)
                    if r == 0:
                        for run in t.cell(r, c).paragraphs[0].runs:
                            run.bold = True
    d.save(path)


def build_xlsx(doc: dict, path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name, rows in doc["sheets"].items():
        ws = wb.create_sheet(sheet_name)
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DCE6F1")
        for col in ws.columns:
            width = max(len(str(c.value)) for c in col) + 2
            ws.column_dimensions[col[0].column_letter].width = min(width, 60)
    wb.properties.title = doc["title"]
    wb.save(path)


def main() -> None:
    builders = {".pdf": build_pdf, ".docx": build_docx, ".xlsx": build_xlsx}
    with REGISTRY.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=REGISTRY_FIELDS)
        writer.writeheader()
        for doc in DOCUMENTS:
            path = RAW / doc["file"]
            path.parent.mkdir(parents=True, exist_ok=True)
            builders[path.suffix](doc, path)
            writer.writerow({k: doc[k] for k in REGISTRY_FIELDS})
            print(f"created {path.relative_to(ROOT)}")
    print(f"\n{len(DOCUMENTS)} documents written, registry at {REGISTRY.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
